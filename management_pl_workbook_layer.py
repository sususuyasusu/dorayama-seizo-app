#!/usr/bin/env python3
"""管理会計PLのコスト分析用読み取り層。

Excelへは書き込まず、保存済みの値と科目構造だけを画面へ渡す。
"""
import os
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

import openpyxl
import openpyxl.utils


# 原本は2026-08-16以降、iCloud Downloadsから運用資料/予算実績表へ移動済み。
# 読み取り専用。存在しない環境（Render等）ではJSONスナップショットへ自動フォールバック。
DEFAULT_PATH = Path(
    "/Users/suzuki3/Library/CloudStorage/Dropbox-Detale/D& W/どら山/運用資料/予算実績表/"
    "【第10期 どら山】管理会計PL.xlsx"
)
WORKBOOK_PATH = Path(os.environ.get("DORAYAMA_MANAGEMENT_PL_PATH", DEFAULT_PATH))
SHEET_NAME = "01_月次PL（管理会計）"
MONTH_COLUMNS = [
    ("B", "2月", "2026/2", "確定"),
    ("C", "3月", "2026/3", "確定"),
    ("D", "4月", "2026/4", "確定"),
    ("E", "5月", "2026/5", "確定"),
    ("F", "6月", "2026/6", "確定"),
    ("G", "7月", "2026/7", "確定"),
    ("H", "8月", "2026/8", "進行中"),
    ("I", "9月", "2026/9", "未入力"),
    ("J", "10月", "2026/10", "未入力"),
    ("K", "11月", "2026/11", "未入力"),
    ("L", "12月", "2026/12", "未入力"),
    ("M", "1月（翌年）", "2027/1", "未入力"),
]
SECTIONS = [
    {"id": "sales", "label": "売上", "rows": range(5, 8)},
    {"id": "materials", "label": "材料・包材原価／粗利益", "rows": range(8, 17)},
    {"id": "labor", "label": "人件費", "rows": range(17, 25)},
    {"id": "other", "label": "その他コスト", "rows": range(25, 45)},
    {"id": "profit", "label": "利益", "rows": range(45, 50)},
]
_CACHE = {"key": None, "value": None}


def _amount(value):
    if value in (None, ""):
        return None
    try:
        return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, ValueError, TypeError):
        return None


def _rate(value):
    if value in (None, ""):
        return None
    try:
        return round(float(value) * 100, 1)
    except (TypeError, ValueError):
        return None


def _cell_value(sheet, cell, is_rate=False):
    value = sheet[cell].value
    return _rate(value) if is_rate else _amount(value)


RATE_LABELS = {"按分に使った本店売上シェア", "粗利益率", "労働分配率（人件費÷粗利益）", "営業利益÷粗利益"}
PROFIT_LABELS = {"売上総利益（粗利益）", "営業利益", "経常利益"}
TOTAL_LABELS = {"売上合計", "材料・包材原価 計", "人件費 合計", "販売管理費 計"} | PROFIT_LABELS
# この科目の行を過ぎたら次の区分へ進む（売上→原価・粗利益→人件費→その他コスト→利益）
SECTION_BREAKS = ["売上合計", "粗利益率", "労働分配率（人件費÷粗利益）", "販売管理費 計"]
HEADER_ROW = 4
FIRST_ROW, LAST_ROW = 5, 120


def _line_kind(label):
    if label in RATE_LABELS:
        return "rate"
    if label in PROFIT_LABELS:
        return "profit"
    if label in TOTAL_LABELS:
        return "total"
    return "detail"


def _month_status(header, has_values, default):
    text = str(header or "")
    if "暫定" in text:
        return "暫定"
    if "まで" in text or "進行中" in text:
        return "進行中"
    if header in (None, ""):
        return default
    return "確定" if has_values else "未入力"


def read_workbook(path):
    path = Path(path)
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET_NAME not in book.sheetnames:
        raise ValueError(f"必要なシートがありません: {SHEET_NAME}")
    sheet = book[SHEET_NAME]
    grid = {}
    for r_index, row_cells in enumerate(sheet.iter_rows(min_row=1, max_row=LAST_ROW, max_col=15), start=1):
        for c_index, cell in enumerate(row_cells, start=1):
            if cell.value not in (None, ""):
                grid[(r_index, c_index)] = cell.value

    def cell(column, row):
        return grid.get((row, openpyxl.utils.column_index_from_string(column)))

    lines = []
    section_index = 0
    for row in range(FIRST_ROW, LAST_ROW + 1):
        raw = cell("A", row)
        if raw in (None, ""):
            continue
        label = str(raw).strip()
        section = SECTIONS[min(section_index, len(SECTIONS) - 1)]
        is_rate = label in RATE_LABELS
        lines.append({
            "row": row,
            "label": label,
            "section": section["id"],
            "sectionLabel": section["label"],
            "kind": _line_kind(label),
            "format": "rate" if is_rate else "amount",
            "values": {
                key: (_rate if is_rate else _amount)(cell(column, row))
                for column, key, _label, _status in MONTH_COLUMNS
            },
            "confirmedTotal": (_rate if is_rate else _amount)(cell("N", row)),
            "previousPeriod": (_rate if is_rate else _amount)(cell("O", row)),
        })
        if section_index < len(SECTION_BREAKS) and label == SECTION_BREAKS[section_index]:
            section_index += 1
    by_label = {}
    for line in lines:
        by_label.setdefault(line["label"], line)

    def value(label, month):
        return by_label.get(label, {}).get("values", {}).get(month)

    months = []
    for column, key, label, status in MONTH_COLUMNS:
        has_values = value("売上合計", key) is not None
        months.append({
            "column": column, "key": key, "label": label,
            "status": _month_status(cell(column, HEADER_ROW), has_values, status),
        })

    series = []
    for month in months:
        key = month["key"]
        sales = value("売上合計", key)
        if sales is None:
            continue
        total_labor = value("人件費 合計", key)
        event_staffing = value("マネキン費（催事販売員の派遣）", key)
        other_cost = value("販売管理費 計", key)
        series.append({
            **month,
            "sales": sales,
            "storeSales": value("店舗売上（本店）", key),
            "eventSales": value("催事売上", key),
            "internalLabor": total_labor - event_staffing
            if total_labor is not None and event_staffing is not None else None,
            "eventStaffing": event_staffing,
            "labor": total_labor,
            "rent": value("地代家賃", key),
            "material": value("材料・包材原価 計", key),
            "other": other_cost,
            "sellingExpenses": total_labor + other_cost
            if other_cost is not None and total_labor is not None else None,
            "operatingProfit": value("営業利益", key),
            "nonOperating": value("雑収入ほか（営業外）", key),
            "profit": value("経常利益", key),
        })
    confirmed = [row for row in series if row["status"] == "確定"]

    def total(field):
        values = [row[field] for row in confirmed if row.get(field) is not None]
        return sum(values) if values else None

    period = (
        f"{confirmed[0]['key']}〜{confirmed[-1]['key']}確定".replace("月〜", "〜")
        if confirmed else "確定月なし"
    )
    stat = path.stat()
    return {
        "available": True,
        "readOnly": True,
        "fileName": path.name,
        "sourceSheet": SHEET_NAME,
        "updatedAt": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="minutes"),
        "months": months,
        "sections": [{"id": section["id"], "label": section["label"]} for section in SECTIONS],
        "lines": lines,
        "series": series,
        "summary": {
            "period": period,
            "sales": total("sales"),
            "internalLabor": total("internalLabor"),
            "eventStaffing": total("eventStaffing"),
            "labor": total("labor"),
            "material": total("material"),
            "other": total("other"),
            "sellingExpenses": total("sellingExpenses"),
            "totalCost": (
                total("material") + total("sellingExpenses")
                if total("material") is not None and total("sellingExpenses") is not None else None
            ),
            "operatingProfit": total("operatingProfit"),
            "nonOperating": total("nonOperating"),
            "profit": total("profit"),
        },
        "rule": "Excelを正本として読み取り。催事販売員は内部人件費から除外し、外注費として別表示。",
    }


SNAPSHOT_FALLBACK_PATH = Path(__file__).resolve().parent / "data" / "management_pl_workbook_snapshot.json"


def _snapshot_fallback():
    """Excel原本を読めない環境（Render等）ではリポジトリ同梱スナップショットを使う。

    スナップショットは tools/build_management_pl_snapshot.py がExcelから読み取り専用で生成。
    """
    import json
    data = json.loads(SNAPSHOT_FALLBACK_PATH.read_text(encoding="utf-8"))
    data["source"] = "snapshot"
    data["rule"] = (data.get("rule") or "") + "（Excel原本を直接読めない環境のため、保存済みスナップショットを表示）"
    return data


def get_cost_analysis():
    try:
        key = (str(WORKBOOK_PATH), WORKBOOK_PATH.stat().st_mtime_ns)
        if _CACHE["key"] != key:
            _CACHE["value"] = read_workbook(WORKBOOK_PATH)
            _CACHE["key"] = key
        return _CACHE["value"]
    except (FileNotFoundError, OSError, ValueError, KeyError) as error:
        try:
            key = ("snapshot", SNAPSHOT_FALLBACK_PATH.stat().st_mtime_ns)
            if _CACHE["key"] != key:
                _CACHE["value"] = _snapshot_fallback()
                _CACHE["key"] = key
            return _CACHE["value"]
        except (FileNotFoundError, OSError, ValueError, KeyError):
            pass
        return {
            "available": False,
            "readOnly": True,
            "fileName": WORKBOOK_PATH.name,
            "sourceSheet": SHEET_NAME,
            "updatedAt": None,
            "months": [],
            "sections": [],
            "lines": [],
            "series": [],
            "summary": {},
            "rule": "管理会計PLを取得できないため、既存の確定集計を表示。",
            "reason": str(error),
        }
